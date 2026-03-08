import pandas as pd
import numpy as np
import os
import re
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill

file_path = '/chikiet/kata2025/ragketoan/BC_QUYET_TOAN_HOANGHUYPHAT_2023.xlsx'
output_dir = '/chikiet/kata2025/ragketoan/docs/hoang-huy-phat'

def clean_name(name):
    if not isinstance(name, str): return 'UNKNOWN'
    name = ' '.join(name.strip().upper().split())
    name = name.replace('(KM)', '').replace('KM ', '').strip()
    return name

def group_item(name):
    if 'NABATI' in name: return 'NBT - BÁNH NABATI'
    if 'VISSAN' in name or 'VS ' in name or 'XÚC XÍCH' in name: return 'VS - XÚC XÍCH VISSAN'
    if 'SANEST' in name or 'SANVINEST' in name or 'NƯỚC YẾN' in name or 'YẾN SÀO' in name: return 'SN - YẾN SÀO/NƯỚC YẾN'
    if 'CHOLIMEX' in name or 'CLM' in name: return 'CLM - GIA VỊ CHOLIMEX'
    if 'ANLENE' in name or 'AL ' in name: return 'AL - SỮA ANLENE'
    if 'XĂNG' in name or 'RON95' in name: return 'FUEL - XĂNG'
    if 'DẦU' in name or 'DO 0.05S' in name: return 'FUEL - DẦU DO'
    if 'TƯƠNG ỚT' in name: return 'CLM - TƯƠNG ỚT (CHUNG)'
    if 'PHÔ MAI' in name or 'BEL ' in name: return 'BEL - PHÔ MAI'
    if 'CHIẾT KHẤU' in name: return 'SERVICE - CHIẾT KHẤU'
    if 'NƯỚC KHOÁNG' in name or 'SANNA' in name: return 'SN - NƯỚC KHOÁNG SANNA'
    if 'ÁO' in name or 'QUẦN' in name or 'BURBERRY' in name: return 'FASHION - QUẦN ÁO/PHỤ KIỆN'
    if 'GIÀY' in name or 'LOAFERS' in name: return 'FASHION - GIÀY DÉP'
    if 'DÂY LƯNG' in name or 'KHÓA LƯNG' in name: return 'FASHION - THẮT LƯNG'
    if 'BÁNH' in name: return 'FOOD - BÁNH (CHUNG)'
    if 'KẸO' in name: return 'FOOD - KẸO (CHUNG)'
    name = re.sub(r'(\d+)\s?(G|ML|KG|ML|HỘP|LON|THÙNG|LỌ|GÓI|CHAI|HŨ|X)', '', name)
    name = ' '.join([w for w in name.split() if len(w) > 2])
    return name[:50].strip()

def process_hhp_xnt_styled():
    print("Reading Excel...")
    inv_df = pd.read_excel(file_path, sheet_name='Ke khai Hoa don 2023')
    xnt_init_df = pd.read_excel(file_path, sheet_name='Tong hop XNT 2023')
    
    inv_df['Tên hàng cleaned'] = inv_df['Tên hàng'].apply(clean_name)
    inv_df['Grouped Name'] = inv_df['Tên hàng cleaned'].apply(group_item)
    
    # Apply threshold for groups to stay under 300
    counts = inv_df['Grouped Name'].value_counts()
    top_groups = counts.index[:250]
    inv_df['Grouped Name'] = inv_df['Grouped Name'].apply(lambda x: x if x in top_groups else 'OTHER - CÁC MẶT HÀNG KHÁC')
    
    # Calculate Nhập/Xuất per group
    inv_2023 = inv_df[(inv_df['Ngày'] >= '2023-01-01') & (inv_df['Ngày'] <= '2023-12-31')].copy()
    
    summary = inv_2023.groupby(['Grouped Name', 'Loại']).agg({
        'Số lượng': 'sum',
        'Thành tiền': 'sum'
    }).unstack(fill_value=0)
    
    # Flatten columns
    summary.columns = [f"{col[0]}_{col[1]}" for col in summary.columns]
    for col in ['Số lượng_Mua vào', 'Số lượng_Bán ra', 'Thành tiền_Mua vào', 'Thành tiền_Bán ra']:
        if col not in summary.columns: summary[col] = 0
        
    # Get ĐVT and Opening Stock from XNT sheet
    xnt_init_df['Tên hàng cleaned'] = xnt_init_df['Tên hàng'].apply(clean_name)
    xnt_init_df['Grouped Name'] = xnt_init_df['Tên hàng cleaned'].apply(group_item)
    xnt_init_df['Grouped Name'] = xnt_init_df['Grouped Name'].apply(lambda x: x if x in top_groups else 'OTHER - CÁC MẶT HÀNG KHÁC')
    
    opening = xnt_init_df.groupby('Grouped Name').agg({
        'ĐVT': 'first',
        'Tồn đầu 01/01/2023': 'sum'
    })
    
    # Merge
    final_df = opening.join(summary, how='outer').fillna(0)
    
    # Price estimation for opening/closing
    # Avg Purchase Price = Nhập Tiền / Nhập SL
    def get_avg_price(row):
        if row['Số lượng_Mua vào'] > 0:
            return row['Thành tiền_Mua vào'] / row['Số lượng_Mua vào']
        elif row['Số lượng_Bán ra'] > 0:
            return (row['Thành tiền_Bán ra'] / row['Số lượng_Bán ra']) * 0.8 # Assume 20% margin
        return 0

    final_df['Avg Price'] = final_df.apply(get_avg_price, axis=1)
    
    # Tồn đầu Tiền
    final_df['Tồn Đầu_Tiền'] = final_df['Tồn đầu 01/01/2023'] * final_df['Avg Price']
    
    # Tồn Cuối SL = Đầu + Nhập - Xuất
    final_df['Tồn Cuối_SL'] = final_df['Tồn đầu 01/01/2023'] + final_df['Số lượng_Mua vào'] - final_df['Số lượng_Bán ra']
    # Tồn Cuối Tiền = Đầu Tiền + Nhập Tiền - (Xuất SL * Avg Price) or proportional
    # Using Accounting formula: Value = End SL * Avg Price
    final_df['Tồn Cuối_Tiền'] = final_df['Tồn Cuối_SL'] * final_df['Avg Price']

    # --- Excel generation ---
    wb = Workbook()
    ws = wb.active
    ws.title = "XNT 2023 Grouped"
    
    # Headers
    header_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    font_bold = Font(bold=True)
    center_align = Alignment(horizontal="center", vertical="center")
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    
    # Row 1: Merged Headers
    ws.merge_cells('A1:A2')
    ws['A1'] = 'MẶT HÀNG / VT'
    ws.merge_cells('B1:B2')
    ws['B1'] = 'ĐVT'
    
    ws.merge_cells('C1:D1')
    ws['C1'] = 'TỒN ĐẦU KỲ'
    ws.merge_cells('E1:F1')
    ws['E1'] = 'NHẬP TRONG KỲ'
    ws.merge_cells('G1:H1')
    ws['G1'] = 'XUẤT TRONG KỲ'
    ws.merge_cells('I1:J1')
    ws['I1'] = 'TỒN CUỐI KỲ'
    
    # Row 2: Sub-headers
    sub_headers = ['SL', 'TIỀN', 'SL', 'TIỀN', 'SL', 'TIỀN', 'SL', 'TIỀN']
    for i, sh in enumerate(sub_headers):
        cell = ws.cell(row=2, column=3+i)
        cell.value = sh
        
    # Styles for headers
    for r in [1, 2]:
        for c in range(1, 11):
            cell = ws.cell(row=r, column=c)
            cell.fill = header_fill
            cell.font = font_bold
            cell.alignment = center_align
            cell.border = thin_border

    # Data rows
    row_idx = 3
    for name, row in final_df.iterrows():
        # A: Name, B: Unit
        ws.cell(row=row_idx, column=1, value=name)
        ws.cell(row=row_idx, column=2, value=row['ĐVT'] if row['ĐVT'] != 0 else 'Cái')
        
        # C, D: Tồn đầu
        ws.cell(row=row_idx, column=3, value=row['Tồn đầu 01/01/2023'])
        ws.cell(row=row_idx, column=4, value=row['Tồn Đầu_Tiền'])
        
        # E, F: Nhập
        ws.cell(row=row_idx, column=5, value=row['Số lượng_Mua vào'])
        ws.cell(row=row_idx, column=6, value=row['Thành tiền_Mua vào'])
        
        # G, H: Xuất
        ws.cell(row=row_idx, column=7, value=row['Số lượng_Bán ra'])
        ws.cell(row=row_idx, column=8, value=row['Thành tiền_Bán ra'])
        
        # I, J: Tồn cuối
        ws.cell(row=row_idx, column=9, value=row['Tồn Cuối_SL'])
        ws.cell(row=row_idx, column=10, value=row['Tồn Cuối_Tiền'])
        
        for c in range(1, 11):
            ws.cell(row=row_idx, column=c).border = thin_border
            if c >= 3:
                ws.cell(row=row_idx, column=c).number_format = '#,##0'
        
        row_idx += 1

    # Adjust column widths
    ws.column_dimensions['A'].width = 50
    ws.column_dimensions['B'].width = 10
    for c in ['C', 'D', 'E', 'F', 'G', 'H', 'I', 'J']:
        ws.column_dimensions[c].width = 15

    output_file = os.path.join(output_dir, 'Bao_cao_XNT_HHP_2023.xlsx')
    wb.save(output_file)
    print(f"Success! Styled XNT report saved in {output_file}")

if __name__ == '__main__':
    process_hhp_xnt_styled()
