import pandas as pd
from openpyxl import load_workbook

def update_excel():
    file_path = '/chikiet/kata2025/ragketoan/docs/huyvu/sosach/nam2024/TK 331 (3).xlsx'
    
    # Read the data from 'sai' sheet
    df_sai = pd.read_excel(file_path, sheet_name='sai')
    
    # Load the workbook to update
    book = load_workbook(file_path)
    
    # Check if target sheet exists
    target_sheet_name = 'TK 331 NĂM 2024 CHƯA CÓ PS NỢ'
    if target_sheet_name not in book.sheetnames:
        print(f"Sheet {target_sheet_name} not found!")
        return

    sheet = book[target_sheet_name]
    
    # We want to update the values. 
    # Since the structures are identical (127 rows, same suppliers), 
    # we can iterate and update.
    
    # Get headers from excel to find column indices
    headers = [cell.value for cell in sheet[1]]
    col_map = {col: i+1 for i, col in enumerate(headers)}
    
    # Ensure columns match
    required_cols = ['Tên nhà cung cấp', 'Số dư đầu kỳ', 'Phát sinh Nợ (Trả)', 'Phát sinh Có (Mua)', 'Số dư cuối kỳ']
    for col in required_cols:
        if col not in col_map:
            print(f"Column {col} not found in target sheet!")
            return

    # Update the data
    # We match by 'Tên nhà cung cấp' just to be safe
    sai_data = df_sai.set_index('Tên nhà cung cấp').to_dict('index')
    
    updated_count = 0
    # Start from row 2 (row 1 is header)
    for row_idx in range(2, sheet.max_row + 1):
        supplier_name = sheet.cell(row=row_idx, column=col_map['Tên nhà cung cấp']).value
        if supplier_name in sai_data:
            row_data = sai_data[supplier_name]
            sheet.cell(row=row_idx, column=col_map['Số dư đầu kỳ']).value = row_data['Số dư đầu kỳ']
            sheet.cell(row=row_idx, column=col_map['Phát sinh Nợ (Trả)']).value = row_data['Phát sinh Nợ (Trả)']
            sheet.cell(row=row_idx, column=col_map['Phát sinh Có (Mua)']).value = row_data['Phát sinh Có (Mua)']
            sheet.cell(row=row_idx, column=col_map['Số dư cuối kỳ']).value = row_data['Số dư cuối kỳ']
            updated_count += 1
            
    book.save(file_path)
    print(f"Updated {updated_count} rows in sheet {target_sheet_name}")

if __name__ == "__main__":
    update_excel()
