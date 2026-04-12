import openpyxl
import os

os.environ['TMPDIR'] = '/chikiet/kata2025/ragketoan/tmp'
file_path = '/chikiet/kata2025/ragketoan/docs/hoang-huy-phat/XNT_HHP_2023.xlsx'
wb = openpyxl.load_workbook(file_path, data_only=True)

issues = []
for sheetname in wb.sheetnames:
    if 'Tháng' not in sheetname: continue
    ws = wb[sheetname]
    for r in range(2, ws.max_row):
        stt = ws.cell(row=r, column=1).value
        ma = ws.cell(row=r, column=2).value
        name = ws.cell(row=r, column=3).value
        if stt is not None and isinstance(stt, (int, float)):
            # Check Nhap
            if (ws.cell(row=r, column=6).value or 0) > 0 and (ws.cell(row=r, column=7).value or 0) == 0:
                issues.append(f"{sheetname} | {ma} | {name} | Col 7 (Nhap VNĐ is 0)")
            # Check Giá Vốn
            if (ws.cell(row=r, column=8).value or 0) > 0 and (ws.cell(row=r, column=11).value or 0) == 0:
                issues.append(f"{sheetname} | {ma} | {name} | Col 11 (Gia Von is 0)")
            # Check Tồn Cuối
            if (ws.cell(row=r, column=12).value or 0) > 0 and (ws.cell(row=r, column=13).value or 0) == 0:
                issues.append(f"{sheetname} | {ma} | {name} | Col 13 (Ton Cuoi is 0)")

print(f"Total issues found: {len(issues)}")
for issue in issues[:20]:
    print(issue)
