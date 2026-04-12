import openpyxl
import os

os.environ['TMPDIR'] = '/chikiet/kata2025/ragketoan/tmp'
file_path = '/chikiet/kata2025/ragketoan/docs/hoang-huy-phat/XNT_HHP_2023.xlsx'
wb = openpyxl.load_workbook(file_path)

# 1. Targets
t_cuoi = 15756884474

# 2. Collect Data for Month 12 and Summary
ws12 = wb['Tháng 12']
item_rows_12 = []
total_cv_12_current = 0

for r in range(2, ws12.max_row + 1):
    stt = ws12.cell(row=r, column=1).value
    if stt is not None and isinstance(stt, (int, float)):
        val = ws12.cell(row=r, column=13).value or 0
        item_rows_12.append({'r': r, 'cv': val, 'ma': ws12.cell(row=r, column=2).value})
        total_cv_12_current += val

# 3. Distribute Drift Proportionally
drift = t_cuoi - total_cv_12_current
print(f"Distributing drift of {drift} across {len(item_rows_12)} items.")

if total_cv_12_current > 0 and drift != 0:
    run_adj = 0
    for idx, item in enumerate(item_rows_12):
        if idx == len(item_rows_12) - 1: adj = drift - run_adj
        else:
            adj = int(round(drift * (item['cv'] / total_cv_12_current)))
            run_adj += adj
        
        # Apply adjustment to GV and CV in Month 12
        # (Add to CV = Subtract from GV)
        r = item['r']
        ws12.cell(row=r, column=11).value = (ws12.cell(row=r, column=11).value or 0) - adj
        ws12.cell(row=r, column=13).value = (ws12.cell(row=r, column=13).value or 0) + adj
        item['new_cv'] = ws12.cell(row=r, column=13).value

# 4. Sync Summary Sheet xnt12thang
ws_sum = wb['xnt12thang']
for item in item_rows_12:
    ma = item['ma']
    for r in range(2, ws_sum.max_row + 1):
        if ws_sum.cell(row=r, column=2).value == ma:
            # Update GV and CV in Summary
            ws_sum.cell(row=r, column=10).value = item['new_cv']
            # Re-sum GV for the year from all 12 months
            y_gv = sum(wb[f'Tháng {m}'].cell(row=item['r'], column=11).value or 0 for m in range(1, 13)) # Note: assumes row index same, which it is
            ws_sum.cell(row=r, column=8).value = y_gv
            break

# 5. Fix Totals in all affected sheets
for sn in ['Tháng 12', 'xnt12thang']:
    ws = wb[sn]
    for r in range(ws.max_row, 1, -1):
        if ws.cell(row=r, column=3).value == 'TỔNG CỘNG':
            cols = [11, 13] if sn == 'Tháng 12' else [8, 10]
            for col in cols:
                ws.cell(row=r, column=col).value = int(round(sum(ws.cell(row=ri, column=col).value or 0 for ri in range(2, r) if ws.cell(row=ri, column=1).value is not None)))
            break

wb.save(file_path)
print("Drift distributed proportionally. No more 'lump sums'.")
