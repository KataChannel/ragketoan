import openpyxl
import pandas as pd

file_path = '/chikiet/kata2025/ragketoan/docs/hoang-huy-phat/XNT_HHP_2023.xlsx'
wb = openpyxl.load_workbook(file_path, data_only=True)

print("--- VERIFICATION REPORT ---")

# 1. Hoadon Sheet
ws_h = wb['Hoadon']
m1_n = ws_h.cell(row=2, column=5).value
m1_x = ws_h.cell(row=3, column=5).value
print(f"Hoadon M1 Nhập: {m1_n} (Target: 14226694560)")
print(f"Hoadon M1 Xuất: {m1_x} (Target: 11386534989)")

# 2. Tháng 12 Total Closing
ws12 = wb['Tháng 12']
total_r12 = ws12.max_row
for r in range(ws12.max_row, 1, -1):
    if ws12.cell(row=r, column=3).value == 'TỔNG CỘNG':
        total_r12 = r; break
closing_12 = ws12.cell(row=total_r12, column=13).value
print(f"Tháng 12 Total Closing: {closing_12} (Target: 15756884474)")

# 3. Negatives Check
negatives = []
for sheet in wb.sheetnames:
    ws = wb[sheet]
    for r in range(2, ws.max_row + 1):
        for c in [5, 7, 8, 9, 10, 11, 13]:
            try:
                val = ws.cell(row=r, column=c).value
                if isinstance(val, (int, float)) and val < 0:
                    negatives.append(f"{sheet}!R{r}C{c}: {val}")
            except: pass

if not negatives:
    print("Zero Negative Numbers Found! [PASS]")
else:
    print(f"Found {len(negatives)} Negative Numbers! [FAIL]")
    print(negatives[:5])

# 4. Summary Total Closing
ws_sum = wb['xnt12thang']
for r in range(ws_sum.max_row, 1, -1):
    if ws_sum.cell(row=r, column=3).value == 'TỔNG CỘNG':
        closing_sum = ws_sum.cell(row=r, column=10).value
        print(f"Summary Total Closing: {closing_sum} (Target: 15756884474)")
        break

if closing_sum == closing_12 == 15756884474:
    print("Closing Balance Sync: [PASS]")
else:
    print("Closing Balance Sync: [FAIL]")
