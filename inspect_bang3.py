import openpyxl

file_path = '/chikiet/kata2025/ragketoan/BC_QUYET_TOAN_HOANGHUYPHAT_2023.xlsx'
wb = openpyxl.load_workbook(file_path, read_only=True)

found = False
for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=500, values_only=True), 1):
        for col_idx, val in enumerate(row):
            if val is not None:
                v_str = str(val).lower().strip()
                if 'bảng' in v_str or 'tháng' in v_str or 'bán' in v_str or 'mua' in v_str:
                    print(f"[{sheet_name}] Row {row_idx}, Col {col_idx}: {v_str}")
                    if 'bảng' in v_str:
                        found = True
                    break

wb.close()
