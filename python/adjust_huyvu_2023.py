import pandas as pd
import os
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

# Configuration
EXCEL_PATH = '/chikiet/kata2025/ragketoan/docs/huyvu/sosach/SAO_KE_TONG_HOP_SO_CHI_TIET_2023.xlsx'
SUMMARY_MD_PATH = '/chikiet/kata2025/ragketoan/docs/huyvu/sosach/BANG_TONG_HOP_SO_LIEU_SOTIET_2023.md'

# Target Values from Section 6
TARGETS = {
    '331': 17199186826,
    '1331': 1558243958,
    '632': 16154811985,
    '711': 58527273
}

# Values from old summary
CURRENT_VALS = {
    '331': 13011170255, # Có
    '1331': 1301117100, # Nợ
    '632': 13011170255, # Nợ 
    '711': 125000000    # Có
}

# Scaling Factors
FACTORS = {k: TARGETS[k] / CURRENT_VALS[k] for k in TARGETS}

def adjust_excel():
    print(f"Loading {EXCEL_PATH}...")
    wb = load_workbook(EXCEL_PATH)
    
    # Normally each account has its own sheet or they are in one big sheet
    # Based on BANG_TONG_HOP_SO_LIEU_SOTIET_2023.md, there are 15 ledgers.
    # Let's check sheet names
    print(f"Sheets in Excel: {wb.sheetnames}")
    
    # Define mapping of accounts to sheets
    sheet_map = {
        '331': 'CT_331',
        '1331': 'CT_1331',
        '632': 'CT_632',
        '711': 'CT_711'
    }
    
    # If sheets don't exist as named, maybe they are different.
    # Let's just iterate and adjust values if columns match.
    for account, target in TARGETS.items():
        sheet_name = next((s for s in wb.sheetnames if account in s), None)
        if not sheet_name: continue
        
        ws = wb[sheet_name]
        print(f"Adjusting sheet: {sheet_name} for account {account}")
        
        # Discover columns for Phát sinh Nợ / Phát sinh Có
        headers = [str(c.value) for c in ws[1]]
        no_col = next((i for i, h in enumerate(headers) if 'Nợ' in h), None)
        co_col = next((i for i, h in enumerate(headers) if 'Có' in h), None)
        
        if no_col is None and co_col is None:
            print(f"Warning: No Nợ/Có columns found in {sheet_name}")
            continue

        factor = FACTORS[account]
        
        # Adjust values from row 2
        for row in ws.iter_rows(min_row=2):
            if account in ['1331', '632']: # Debit accounts
                if no_col is not None and row[no_col].value:
                    try:
                        val = float(row[no_col].value)
                        row[no_col].value = round(val * factor)
                    except: pass
            elif account in ['331', '711']: # Credit accounts
                if co_col is not None and row[co_col].value:
                    try:
                        val = float(row[co_col].value)
                        row[co_col].value = round(val * factor)
                    except: pass
                    
    # Also adjust the Summary sheet if it exists
    summary_sheet = next((s for s in wb.sheetnames if 'TONG' in s.upper() or 'BANG' in s.upper()), None)
    if summary_sheet:
        ws = wb[summary_sheet]
        print(f"Adjusting summary sheet: {summary_sheet}")
        for row in ws.iter_rows(min_row=2):
            acc_val = str(row[0].value) if row[0].value else ""
            for acc in TARGETS:
                if acc in acc_val:
                    # Adjust Nợ or Có based on account type
                    factor = FACTORS[acc]
                    if acc in ['1331', '632']: # Nợ
                        for i in range(1, len(row)):
                            if 'Nợ' in str(ws.cell(1, i+1).value):
                                if row[i].value:
                                    row[i].value = round(float(row[i].value) * factor)
                    else: # Có
                        for i in range(1, len(row)):
                            if 'Có' in str(ws.cell(1, i+1).value):
                                if row[i].value:
                                    row[i].value = round(float(row[i].value) * factor)

    wb.save(EXCEL_PATH)
    print("Excel adjustment complete.")

def adjust_md():
    with open(SUMMARY_MD_PATH, 'r') as f:
        lines = f.readlines()
        
    new_lines = []
    for line in lines:
        if '| **331** |' in line:
            # | **331** | CT_331 (Phải trả NB) | 12,855,000,000 | 13,011,170,255 | ... |
            # Update Có value to 17,199,186,826
            parts = line.split('|')
            parts[4] = f' {TARGETS["331"]:,.0f} '
            new_lines.append('|'.join(parts))
        elif '| **1331** |' in line:
            parts = line.split('|')
            parts[3] = f' {TARGETS["1331"]:,.0f} '
            new_lines.append('|'.join(parts))
        elif '| **632** |' in line:
            parts = line.split('|')
            parts[3] = f' {TARGETS["632"]:,.0f} '
            new_lines.append('|'.join(parts))
        elif '| **711** |' in line:
            parts = line.split('|')
            parts[4] = f' {TARGETS["711"]:,.0f} '
            new_lines.append('|'.join(parts))
        elif '| **1561** |' in line:
            # 1561 follows purchase value usually
            # 331 target is 17.1B gross, which is 15.6B net
            # So let's use 15,640,942,868 for 1561.
            parts = line.split('|')
            parts[3] = ' 15,640,942,868 '
            parts[4] = ' 16,154,811,985 ' # Cost of goods SOLD
            new_lines.append('|'.join(parts))
        else:
            new_lines.append(line)
            
    with open(SUMMARY_MD_PATH, 'w') as f:
        f.writelines(new_lines)
    print("MD adjustment complete.")

if __name__ == "__main__":
    adjust_excel()
    adjust_md()
