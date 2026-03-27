import openpyxl

f = "/chikiet/kata2025/ragketoan/docs/huyvu/XNT_HuyVu_2023.xlsx"
wb = openpyxl.load_workbook(f)
total_nhap = 0
for name in wb.sheetnames:
    ws = wb[name]
    month_nhap = 0
    for r in range(2, ws.max_row+1):
        v = ws.cell(r, 7).value  # col G = Nhập (VNĐ)
        if v and isinstance(v, (int, float)):
            month_nhap += v
    print(f"{name}: Nhập VNĐ = {month_nhap:,.0f}")
    total_nhap += month_nhap
print(f"\nTỔNG NHẬP CẢ NĂM: {total_nhap:,.0f}")
print(f"Mục tiêu: 14,910,791,883 - 15,640,942,868")
print(f"Chênh lệch vs 15.6B: {total_nhap - 15_640_942_868:,.0f}")
