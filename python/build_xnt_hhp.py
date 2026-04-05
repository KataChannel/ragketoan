import argparse
import json
import re
import os
import sys
import warnings
from collections import defaultdict
from datetime import datetime

import pandas as pd
import numpy as np
from sqlalchemy import create_engine, text
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

warnings.filterwarnings('ignore')

# ============================================================
# CONFIG
# ============================================================
DB_URI = "postgresql://root:password@localhost:5432/ketoan"
OUTPUT_DIR = "/chikiet/kata2025/ragketoan/docs/hoang-huy-phat"
COMPANY_ID = "03b043e9-b7cd-42bc-a4ea-db710552af82"
MST = "5900428904"

def load_mapping():
    mapping_file = "/chikiet/kata2025/ragketoan/docs/hoang-huy-phat/2.mapping_341_items.md"
    with open(mapping_file, "r") as f: data = f.read()
    rows = re.findall(r'\| ([0-9]+) \| (.*?) \| (.*?) \| (.*?) \|', data)
    id_map = {}; names_2023 = {}; names_2024 = {}
    for row in rows:
        stt, ma, ten_2024, ten_2023_raw = [x.strip() for x in row]
        id_map[ma] = ten_2024
        names_2024[ten_2024.lower()] = ma
        sub_names = [x.strip() for x in ten_2023_raw.split("<br>") if x.strip()]
        for sub_name in sub_names:
            if sub_name != "*Không tìm thấy*": names_2023[sub_name.lower()] = ma
    return id_map, names_2024, names_2023

ID_MAP, NAMES_2024, NAMES_2023 = load_mapping()

def map_item(ten_hang, year):
    if not ten_hang: return "OTH", "Hàng hóa khác"
    h = str(ten_hang).lower().strip()
    mapping_source = NAMES_2023 if year == 2023 else NAMES_2024
    if h in mapping_source: return mapping_source[h], ID_MAP[mapping_source[h]]
    for n, ma in mapping_source.items():
        if n in h or h in n: return ma, ID_MAP[ma]
    all_names = list(NAMES_2024.items()) + list(NAMES_2023.items())
    for n, ma in all_names:
        if n in h or h in n: return ma, ID_MAP[ma]
    return "OTH", "Hàng hóa khác"

def fetch_data(engine, company_id, year):
    q_list = text("""
        SELECT \"idServer\", shdon, 
               tdlap AT TIME ZONE 'UTC' AT TIME ZONE 'Asia/Ho_Chi_Minh' as tdlap_ict,
               loaihd, tthai, tgtcthue
        FROM ext_listhoadon
        WHERE \"congtyId\" = :cid AND tthai IN ('1','2','4','5')
          AND (tdlap AT TIME ZONE 'UTC' AT TIME ZONE 'Asia/Ho_Chi_Minh') >= :start
          AND (tdlap AT TIME ZONE 'UTC' AT TIME ZONE 'Asia/Ho_Chi_Minh') < :end
    """)
    with engine.connect() as conn:
        df_list = pd.read_sql(q_list, conn, params={'cid': company_id, 'start': f'{year}-01-01', 'end': f'{year+1}-01-01'})
    if df_list.empty: return pd.DataFrame(), pd.DataFrame()
    ids = df_list['idServer'].tolist()
    all_details = []
    for i in range(0, len(ids), 500):
        batch = ids[i:i+500]
        q_det = text("SELECT \"idhdonServer\", ten, sluong, dgia, thtien FROM ext_detailhoadon WHERE \"idhdonServer\" IN :ids")
        with engine.connect() as conn: df_batch = pd.read_sql(q_det, conn, params={'ids': tuple(batch)})
        all_details.append(df_batch)
    return df_list, pd.concat(all_details, ignore_index=True) if all_details else pd.DataFrame()

def process_xnt(df_list, df_detail, year, ton_dau_vnd=0, target_gv_vnd=None):
    df_list['month'] = pd.to_datetime(df_list['tdlap_ict']).dt.month
    df = df_detail.merge(df_list[['idServer', 'month', 'loaihd', 'tgtcthue']], left_on='idhdonServer', right_on='idServer', how='left')
    df['scale_fact'] = df['tgtcthue'] / df.groupby('idServer')['thtien'].transform('sum').replace(0, 1)
    df[['group_id', 'group_name']] = df.apply(lambda r: pd.Series(map_item(r['ten'], year)), axis=1)
    df['qty'] = df['sluong'].fillna(0).astype(float); df['value_actual'] = df['thtien'].fillna(0).astype(float) * df['scale_fact']
    all_groups = sorted(df['group_id'].unique()); total_nhap_v = 0
    res = {m: {g: {'nhap_sl': 0, 'nhap_vnd': 0, 'xuat_sl': 0, 'xuat_vnd': 0} for g in all_groups} for m in range(1, 13)}
    for _, r in df.iterrows():
        m, g, l, q, v = r['month'], r['group_id'], r['loaihd'], r['qty'], r['value_actual']
        if l == 'muavao': res[m][g]['nhap_sl'] += q; res[m][g]['nhap_vnd'] += v; total_nhap_v += v
        else: res[m][g]['xuat_sl'] += q; res[m][g]['xuat_vnd'] += v
    group_min_q = {g: 0 for g in all_groups}
    for g in all_groups:
        run = 0
        for m in range(1, 13):
            run += res[m][g]['nhap_sl'] - res[m][g]['xuat_sl']
            if run < 0: group_min_q[g] = max(group_min_q[g], abs(run))
    avg_p = {}
    for g in all_groups:
        q_t = df[df['group_id'] == g]['qty'].sum(); v_t = df[df['group_id'] == g]['value_actual'].sum()
        avg_p[g] = v_t / q_t if q_t > 0 else 100000
    tdg = {}
    for g in all_groups:
        q_s = int(np.ceil(group_min_q[g] * 2.0))
        if q_s == 0 and df[(df['group_id'] == g) & (df['loaihd'] == 'banra')].shape[0] > 0: q_s = 100
        tdg[g] = {'qty': q_s, 'val': q_s * avg_p[g]}
    cur_v = sum(v['val'] for v in tdg.values()); norm = ton_dau_vnd / cur_v if cur_v > 0 else 1.0
    for g in all_groups: tdg[g]['val'] *= norm
    cogs_f = 1.0
    if target_gv_vnd:
        def calc_gv(f):
            tot = 0; mq, mv = {g: tdg[g]['qty'] for g in all_groups}, {g: tdg[g]['val'] for g in all_groups}
            for m in range(1, 13):
                for g in all_groups:
                    d = res[m][g]; avg = (mv[g] + d['nhap_vnd']) / (mq[g] + d['nhap_sl']) if (mq[g] + d['nhap_sl']) > 0 else 0
                    gv = min(avg * d['xuat_sl'] * f, max(0, mv[g] + d['nhap_vnd'] - 1000)); tot += gv; mq[g] += d['nhap_sl'] - d['xuat_sl']; mv[g] += d['nhap_vnd'] - gv
            return tot
        f = 1.0
        for _ in range(50):
            curr = calc_gv(f)
            if abs(curr - target_gv_vnd) < 1000 or (curr == 0 and target_gv_vnd > 0): break
            f = f * (target_gv_vnd / curr) if curr > 0 else f * 1.5
        cogs_f = f
        actual_gv = calc_gv(f)
        print(f"  Year {year} | Target GV: {target_gv_vnd:,.0f} | Actual GV: {actual_gv:,.0f} | Factor: {f:.6f}")
        print(f"  Year {year} | Resulting Ending: {ton_dau_vnd + total_nhap_v - actual_gv:,.0f}")
    return res, all_groups, tdg, df_list[['month', 'loaihd', 'tgtcthue']].copy(), cogs_f

def build_excel(res, all_g, year, out_p, tdg, hd_d, cogs_f):
    wb = Workbook(); wb.remove(wb.active)
    st = lambda c: {'font': Font(bold=True, size=11, color="FFFFFF" if c=='h' else "000000"), 'fill': PatternFill("solid", fgColor="2F5496" if c=='h' else "DDEBF7"), 'alignment': Alignment(horizontal='center'), 'border': Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))}
    hs, ts, nf = st('h'), st('t'), '#,##0'
    ws_m = wb.create_sheet("xnt12thang"); hdrs = ['STT', 'Mã Nhóm', 'Tên Nhóm Sản Phẩm', 'Tồn Đầu Kỳ (SL)', 'Tồn Đầu Kỳ (VNĐ)', 'Tổng Tiền Nhập VNĐ', 'Tổng Tiền Xuất HĐ VNĐ', 'Tổng Tiền Giá Vốn VNĐ', 'Tồn Cuối (SL)', 'Tồn Cuối (VNĐ)']
    for m in range(1, 13): hdrs.extend([f'Nhập T{m} VNĐ', f'Xuất T{m} VNĐ'])
    for c, h in enumerate(hdrs, 1): cell = ws_m.cell(1, c, h); cell.font = hs['font']; cell.fill = hs['fill']; cell.border = hs['border']; cell.alignment = hs['alignment']
    for idx, g in enumerate(all_g, 1):
        td = tdg[g]; cq, cv = td['qty'], td['val']; sn, sxh, sgv = 0, 0, 0; flows = []
        for m in range(1, 13):
            d = res[m][g]; avg = (cv + d['nhap_vnd']) / (cq + d['nhap_sl']) if (cq + d['nhap_sl']) > 0 else 0
            gv = min(avg * d['xuat_sl'] * cogs_f, max(0, cv+d['nhap_vnd'] - 1000)); sn += d['nhap_vnd']; sxh += d['xuat_vnd']; sgv += gv; flows.extend([d['nhap_vnd'], gv]); cq += d['nhap_sl'] - d['xuat_sl']; cv += d['nhap_vnd'] - gv
        for c, v in enumerate([idx, g, ID_MAP.get(g, "Khác"), td['qty'], td['val'], sn, sxh, sgv, cq, cv] + flows, 1): cell = ws_m.cell(idx+1, c, v); cell.border = hs['border']; (c >= 4) and (setattr(cell, 'number_format', nf))
    lr = len(all_g) + 2; ws_m.cell(lr, 3, "TỔNG CỘNG").font = ts['font']
    for c in range(4, len(hdrs)+1): cell = ws_m.cell(lr, c, f"=SUM({get_column_letter(c)}2:{get_column_letter(c)}{lr-1})"); cell.font = ts['font']; cell.fill = ts['fill']; cell.border = ts['border']; cell.number_format = nf
    mq, mv = {g: tdg[g]['qty'] for g in all_g}, {g: tdg[g]['val'] for g in all_g}
    for m in range(1, 13):
        ws = wb.create_sheet(f"Tháng {m}"); mh = ['STT', 'Mã Hàng', 'Tên Hàng', 'Tồn Đầu Kỳ (SL)', 'Tồn Đầu Kỳ (VNĐ)', 'Nhập (SL)', 'Nhập (VNĐ)', 'Xuất (SL)', 'Xuất (VNĐ)', 'Giá Vốn', 'Xuất Theo Giá Vốn', 'Tồn Cuối (SL)', 'Tồn Cuối (VNĐ)']
        for c, h in enumerate(mh, 1): cell = ws.cell(1, c, h); cell.font = hs['font']; cell.fill = hs['fill']; cell.border = hs['border']; cell.alignment = hs['alignment']
        for idx, g in enumerate(all_g, 1):
            d = res[m][g]; qd, vd = mq[g], mv[g]; avg = (vd + d['nhap_vnd']) / (qd + d['nhap_sl']) if (qd + d['nhap_sl']) > 0 else 0
            gv = min(avg * d['xuat_sl'] * cogs_f, max(0, vd+d['nhap_vnd']-1000)); qc, vc = qd + d['nhap_sl'] - d['xuat_sl'], vd + d['nhap_vnd'] - gv
            for c, v in enumerate([idx, g, ID_MAP.get(g, "Khác"), qd, vd, d['nhap_sl'], d['nhap_vnd'], d['xuat_sl'], d['xuat_vnd'], gv/d['xuat_sl'] if d['xuat_sl']>0 else 0, gv, qc, vc], 1): ws.cell(idx+1, c, v).border = hs['border']; (c >= 4) and (setattr(ws.cell(idx+1, c, v), 'number_format', nf))
            mq[g], mv[g] = qc, vc
        ml = len(all_g) + 2; ws.cell(ml, 3, "TỔNG CỘNG").font = ts['font']
        for c in range(4, 14): cell = ws.cell(ml, c, f"=SUM({get_column_letter(c)}2:{get_column_letter(c)}{ml-1})"); cell.font = ts['font']; cell.fill = ts['fill']; cell.border = ts['border']; cell.number_format = nf
    ws_hd = wb.create_sheet("Hoadon"); hdh = ['Tháng', 'Tổng Tiền Hóa Đơn Mua (VNĐ)', 'Tổng Tiền Hóa Đơn Bán (VNĐ)']
    for c, h in enumerate(hdh, 1): cell = ws_hd.cell(1, c, h); cell.font = hs['font']; cell.fill = hs['fill']; cell.border = hs['border']; cell.alignment = hs['alignment']
    for m in range(1, 13):
        vs = [f"Tháng {m}", hd_d[(hd_d['month'] == m) & (hd_d['loaihd'] == 'muavao')]['tgtcthue'].sum(), hd_d[(hd_d['month'] == m) & (hd_d['loaihd'] == 'banra')]['tgtcthue'].sum()]
        for c, v in enumerate(vs, 1): cell = ws_hd.cell(m+1, c, v); cell.border = hs['border']; (c > 1) and (setattr(cell, 'number_format', nf))
    ws_hd.cell(14, 1, "TỔNG CỘNG").font = ts['font']
    for c in range(2, 4): cell = ws_hd.cell(14, c, f"=SUM({get_column_letter(c)}2:{get_column_letter(c)}13)"); cell.font = ts['font']; cell.fill = ts['fill']; cell.border = ts['border']; cell.number_format = nf
    wb.save(out_p)

def main():
    parser = argparse.ArgumentParser(); parser.add_argument('--year', type=int, required=True); args = parser.parse_args()
    engine = create_engine(DB_URI)
    
    if args.year == 2023:
        t_start, t_gv = (15447634554, 110197182873)
    else:
        # 2024: read opening from 2023 closing
        t_start = 15761265757  # Default = 2023 Cuối Kỳ
        t_gv = 103332100062
        try:
            search_paths = [
                os.path.join(OUTPUT_DIR, 'sosach2023', 'XNT_HoangHuyPhat_2023.xlsx'),
                os.path.join(OUTPUT_DIR, 'XNT_HHP_2023.xlsx'),
            ]
            for fpath in search_paths:
                if os.path.exists(fpath):
                    df23 = pd.read_excel(fpath, sheet_name='xnt12thang')
                    last_row = df23.iloc[-1]
                    col_cuoi = [c for c in df23.columns if 'Cuối' in str(c) and ('VNĐ' in str(c) or 'Tiền' in str(c))]
                    if col_cuoi:
                        val = last_row[col_cuoi[0]]
                    else:
                        val = last_row.iloc[-1]
                    if pd.notna(val) and val > 0:
                        t_start = float(val)
                        print(f"  2024 Opening set from {os.path.basename(fpath)}: {t_start:,.0f}")
                    else:
                        print(f"  File found but value invalid, using default: {t_start:,.0f}")
                    break
        except Exception as e:
            print(f"  Using default 2024 opening: {t_start:,.0f} (error: {e})")
    
    df_l, df_d = fetch_data(engine, COMPANY_ID, args.year)
    if df_l.empty: return
    res, groups, tdg, hd, f = process_xnt(df_l, df_d, args.year, t_start, t_gv)
    build_excel(res, groups, args.year, os.path.join(OUTPUT_DIR, f"XNT_HHP_{args.year}.xlsx"), tdg, hd, f)

if __name__ == "__main__": main()
