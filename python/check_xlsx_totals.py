import openpyxl

f = "/chikiet/kata2025/ragketoan/docs/huyvu/XNT_HuyVu_2023.xlsx"
wb = openpyxl.load_workbook(f)
for name in wb.sheetnames:
    ws = wb[name]
    print(f"\n=== Sheet: {name} ===")
    print(f"Rows: {ws.max_row}, Cols: {ws.max_column}")
    # Print first 5 rows
    for r in range(1, min(6, ws.max_row+1)):
        row = [ws.cell(r, c).value for c in range(1, ws.max_column+1)]
        print(row)
    # Find "Tổng Nhập" or sum columns
    for r in range(1, ws.max_row+1):
        for c in range(1, ws.max_column+1):
            v = ws.cell(r, c).value
            if v and isinstance(v, str) and 'nhập' in v.lower():
                print(f"  Found '{v}' at ({r},{c}), next={ws.cell(r,c+1).value}")
