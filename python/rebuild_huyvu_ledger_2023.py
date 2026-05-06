import pandas as pd
import re
import os
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# Config
SOURCE_FILE = '/chikiet/kata2025/ragketoan/docs/huyvu/sosach/nam2023/SỔ SÁCH CÁC TK 2023 HUY VŨ.xlsx'
OUTPUT_DIR = '/chikiet/kata2025/ragketoan/docs/huyvu/sosach/nam2023'
OUTPUT_FILE = os.path.join(OUTPUT_DIR, 'SỔ SÁCH CÁC TK 2023 HUY VŨ FINAL.xlsx')

def extract_supplier_name(desc):
    if not desc: return None
    # Skip generic labels
    if desc in ['SỐ DƯ ĐẦU KỲ', 'TỔNG PHÁT SINH TRONG KỲ', 'SỐ DƯ CUỐI KỲ', 'Thuế GTGT đầu vào', 'Thuế GTGT đầu ra']:
        return None
        
    # Pattern: [Action] - [Supplier]
    match = re.search(r'-\s*(.*)', desc)
    if match: return match.group(1).strip()
    
    # Pattern: [Action] cho/của [Supplier]
    match = re.search(r'(?:cho|của)\s+(.*)', desc, re.IGNORECASE)
    if match: return match.group(1).strip()
    
    return desc.strip()

def process_131(writer):
    print("Processing TK 131...")
    df = pd.read_excel(SOURCE_FILE, sheet_name='131')
    
    # Get Targets from Column J (Unnamed: 9)
    # Row 0: Opening
    opening_bal = df.iloc[0]['Unnamed: 9']
    if pd.isna(opening_bal):
        opening_bal = df.iloc[0]['Đầu kỳ']
        
    # Row -1: Target Closing
    target_closing = df.iloc[-1]['Unnamed: 9']
    
    print(f"  Target Opening: {opening_bal:,.0f}")
    print(f"  Target Closing: {target_closing:,.0f}")
    
    # Filter transaction rows
    mask_trans = ~df['Diễn giải'].isin(['SỐ DƯ ĐẦU KỲ', 'TỔNG PHÁT SINH TRONG KỲ', 'SỐ DƯ CUỐI KỲ'])
    
    # Identify Adjustment Row if exists
    adj_mask = df['Số chứng từ'].astype(str).str.contains('ADJ', na=False)
    
    # Initial calculation without the adjustment row (to see the real gap)
    real_trans_mask = mask_trans & (~adj_mask)
    real_debit = df.loc[real_trans_mask, 'Phát sinh Nợ'].sum()
    real_credit = df.loc[real_trans_mask, 'Phát sinh Có'].sum()
    
    current_gap = (opening_bal + real_debit - real_credit) - target_closing
    print(f"  Calculated Gap (before ADJ): {current_gap:,.0f}")
    
    # Update the adjustment row or create one
    if adj_mask.any():
        idx = df[adj_mask].index[0]
        # If gap > 0, we need more Credit (H)
        if current_gap > 0:
            df.at[idx, 'Phát sinh Có'] = current_gap
            df.at[idx, 'Phát sinh Nợ'] = 0
        else:
            df.at[idx, 'Phát sinh Có'] = 0
            df.at[idx, 'Phát sinh Nợ'] = abs(current_gap)
        df.at[idx, 'Diễn giải'] = f"Điều chỉnh số dư cuối kỳ khớp mục tiêu {target_closing:,.0f}"
    else:
        # Create new row just before totals
        total_idx = df[df['Diễn giải'] == 'TỔNG PHÁT SINH TRONG KỲ'].index[0]
        new_row = {
            'Ngày hạch toán': '31/12/2023',
            'Ngày chứng từ': '31/12/2023',
            'Số chứng từ': 'ADJ_131_BAL',
            'Diễn giải': f"Điều chỉnh số dư cuối kỳ khớp mục tiêu {target_closing:,.0f}",
            'TK Đối ứng': '1111',
            'Đầu kỳ': 0,
            'Phát sinh Nợ': 0 if current_gap > 0 else abs(current_gap),
            'Phát sinh Có': current_gap if current_gap > 0 else 0,
            'Cuối kỳ': 0
        }
        df = pd.concat([df.iloc[:total_idx], pd.DataFrame([new_row]), df.iloc[total_idx:]]).reset_index(drop=True)

    # RECALCULATE RUNNING BALANCE
    current_running = opening_bal
    df.at[0, 'Đầu kỳ'] = opening_bal
    df.at[0, 'Cuối kỳ'] = opening_bal
    
    for i in range(1, len(df)):
        desc = df.at[i, 'Diễn giải']
        if desc in ['TỔNG PHÁT SINH TRONG KỲ', 'SỐ DƯ CUỐI KỲ']:
            continue
        
        nợ = df.at[i, 'Phát sinh Nợ'] or 0
        có = df.at[i, 'Phát sinh Có'] or 0
        
        current_running = current_running + nợ - có
        df.at[i, 'Cuối kỳ'] = current_running
        
    # Update Totals
    total_idx = df[df['Diễn giải'] == 'TỔNG PHÁT SINH TRONG KỲ'].index[0]
    final_idx = df[df['Diễn giải'] == 'SỐ DƯ CUỐI KỲ'].index[0]
    
    sum_nợ = df.loc[mask_trans, 'Phát sinh Nợ'].sum()
    sum_có = df.loc[mask_trans, 'Phát sinh Có'].sum()
    
    df.at[total_idx, 'Phát sinh Nợ'] = sum_nợ
    df.at[total_idx, 'Phát sinh Có'] = sum_có
    df.at[final_idx, 'Cuối kỳ'] = current_running
    
    # Save to Excel
    df.to_excel(writer, sheet_name='131', index=False)

def process_331(writer):
    print("Processing TK 331...")
    df = pd.read_excel(SOURCE_FILE, sheet_name='331')
    
    # Extract Supplier
    df['Supplier'] = df['Diễn giải'].apply(extract_supplier_name)
    
    # Propagation of Supplier to tax/blank rows
    # Group by (Date + Voucher Number)
    df['Key'] = df['Ngày hạch toán'].astype(str) + "_" + df['Số chứng từ'].astype(str)
    
    supplier_map = df[df['Supplier'].notna()].groupby('Key')['Supplier'].first().to_dict()
    
    def fill_supplier(row):
        if pd.notna(row['Supplier']): return row['Supplier']
        return supplier_map.get(row['Key'], "Khác / Chưa phân loại")
        
    df['Supplier'] = df.apply(fill_supplier, axis=1)
    
    # Filter trans rows for summary
    mask_trans = ~df['Diễn giải'].isin(['SỐ DƯ ĐẦU KỲ', 'TỔNG PHÁT SINH TRONG KỲ', 'SỐ DƯ CUỐI KỲ'])
    trans_df = df[mask_trans].copy()
    
    # Create Summary
    summary = trans_df.groupby('Supplier').agg({
        'Phát sinh Nợ': 'sum',
        'Phát sinh Có': 'sum'
    }).reset_index()
    
    opening_331 = df.iloc[0]['Cuối kỳ']
    summary['Dư Cuối Kỳ'] = summary['Phát sinh Có'] - summary['Phát sinh Nợ'] # For 331 usually Credit is positive
    
    # Save main sheet
    df.drop(columns=['Key'], inplace=True)
    df.to_excel(writer, sheet_name='331_Chi_Tiet', index=False)
    
    # Save summary sheet
    summary.to_excel(writer, sheet_name='331_TONG_HOP', index=False)

def finalize_formatting(file_path):
    wb = load_workbook(file_path)
    # Apply accounting format to number columns
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, (int, float)):
                    cell.number_format = '#,##0'
                    
    # Auto-adjust column width
    for ws in wb.worksheets:
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except: pass
            adjusted_width = (max_length + 2)
            ws.column_dimensions[column].width = min(adjusted_width, 60)

    wb.save(file_path)

if __name__ == "__main__":
    with pd.ExcelWriter(OUTPUT_FILE, engine='xlsxwriter') as writer:
        process_131(writer)
        process_331(writer)
        
    finalize_formatting(OUTPUT_FILE)
    print(f"\nDone! Final file saved at: {OUTPUT_FILE}")
