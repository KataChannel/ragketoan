import argparse
import json
from collections import defaultdict
import re
import os
import sys
import pandas as pd
import numpy as np
from sqlalchemy import create_engine, text
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ============================================================
# CONFIG
# ============================================================
DB_URI = os.environ.get("XNT_DB_URI", "postgresql://root:password@localhost:5432/ketoan")
OUTPUT_DIR = "/chikiet/kata2025/ragketoan/docs/huyvu"
COMPANY_ID = "db88c924-206b-4544-9256-c1cd79d417e4" # Huy Vũ
MST = "5900363291"

# ============================================================
# TARGETS & SPECIFIC FILTERS
# ============================================================
OPENING_VALS = {
    2023: 20_528_682_383,
    2024: 19_999_094_250,
    2025: 31_017_722_278,
    2026: 36_468_593_764
}

# Specific SH numbers to skip based on accounting reconciliation (2023 Sales)
SALES_SKIP_2023 = {
    '129', '69', '187', '221', '456', '420', '451', '497', '531', '681', '627', 
    '698', '758', '800', '786', '808', '988', '1010', '964', '1119', '1112', 
    '1123', '1048', '1332', '1249', '1272', '1508', '1463', '1538'
}

# Specific SH numbers to skip based on accounting reconciliation (2023 Purchases)
PURCH_SKIP_2023 = {
    "112762", "114640", "1151", "115328", "12504", "1264", "129748", "130700", "132995", 
    "133902", "1390", "14405", "146770", "146774", "146775", "1477", "1575", "15840", 
    "15905", "16763463", "17575", "177470", "178099", "179887", "1845", "1929293", 
    "194832", "20519", "2064", "20691", "2272", "235925", "255", "265949", "26681", 
    "268", "27678", "2775", "294865", "3090", "3109", "313055", "3171", "3209", 
    "3270", "33062", "35227", "35340", "35341", "354591", "3570", "3647", "3785", 
    "3823", "3972", "4091", "4140", "415467", "427543", "4285", "45189", "45602", 
    "4577", "461536", "471", "4735", "47929", "4908", "500682", "5059", "51554", 
    "51556", "51575", "5275", "52821", "540665", "5495", "5559", "5564", "56188", 
    "5709", "5917", "6010", "606687", "6082", "634869", "6357", "6761", "68601", 
    "68602", "741", "746", "7774", "82589", "82795", "83279", "83284", "833663", 
    "8343", "8424", "8425", "8434", "9888", "99003", "99034"
}

# ============================================================
# MAPPING UTILS
# ============================================================
def get_category_mapping():
    return {
        "PC-DELL-LAT": "Laptop DELL Latitude",
        "PC-DELL-VOS": "Laptop DELL Vostro",
        "PC-DELL-INS": "Laptop DELL Inspiron",
        "PC-DELL-XPS": "Laptop DELL XPS",
        "PC-ASU-VIVO": "Laptop ASUS VivoBook",
        "PC-ASU-ZEN": "Laptop ASUS ZenBook",
        "PC-ASU-EXP": "Laptop ASUS ExpertBook",
        "PC-ASU-ROG": "Laptop ASUS Gaming",
        "PC-LEN-TP": "Laptop Lenovo ThinkPad",
        "PC-LEN-IP": "Laptop Lenovo IdeaPad",
        "PC-LEN-V": "Laptop Lenovo V-Series",
        "PC-MSI-MOD": "Laptop MSI Modern",
        "PC-MSI-GF": "Laptop MSI Gaming",
        "PC-HP-PAV": "Laptop HP Pavilion",
        "PC-HP-PRO": "Laptop HP ProBook",
        "PC-HP-EL": "Laptop HP EliteBook",
        "PC-ACER-ASP": "Laptop ACER Aspire",
        "PC-ACER-NIT": "Laptop ACER Nitro",
        "PC-MAC-AIR": "Apple MacBook Air",
        "PC-MAC-PRO": "Apple MacBook Pro",
        "PC-TAB-IPAD": "Apple iPad",
        "PC-TAB-SAM": "Samsung Galaxy Tab",
        "PC-SYS-I3": "PC Văn phòng Core i3",
        "PC-SYS-I5": "PC Văn phòng Core i5",
        "PC-SYS-I7": "PC Đồ họa Core i7",
        "PC-SYS-G": "PC Văn phòng Entry",
        "PC-WS-XEON": "Workstation / Server",
        "PC-AIO-DELL": "All-in-One DELL",
        "PC-AIO-HP": "All-in-One HP",
        "PC-MINI": "NUC / Mini PC",
        "LCD-SAM-19": "Màn hình Samsung 19-20\"",
        "LCD-SAM-24": "Màn hình Samsung 24-27\"",
        "LCD-DELL-22": "Màn hình DELL 22\"",
        "LCD-DELL-24": "Màn hình DELL 24\"",
        "LCD-DELL-27": "Màn hình DELL 27\"",
        "LCD-ASU-24": "Màn hình ASUS 24\"",
        "LCD-LG-24": "Màn hình LG 24-27\"",
        "LCD-VIEW-24": "Màn hình ViewSonic",
        "LCD-GAM-144": "Màn hình Gaming 144Hz+",
        "LCD-PRO-4K": "Màn hình Đồ họa 4K",
        "PRN-CAN-LBP": "Máy in Laser Canon",
        "PRN-CAN-MF": "Máy in Đa năng Canon",
        "PRN-BRO-HL": "Máy in Laser Brother",
        "PRN-BRO-DCP": "Máy in Đa năng Brother",
        "PRN-HP-LJ": "Máy in Laser HP",
        "PRN-EPS-L": "Máy in Phun Epson",
        "SCN-CAN-CANO": "Máy quét Canon",
        "SCN-HP-SJ": "Máy quét HP",
        "PRN-POS-80": "Máy in Hóa đơn K80",
        "SCN-BAR": "Đầu đọc mã vạch",
        "OFF-CHAIR-ST": "Ghế văn phòng",
        "OFF-DESK-W": "Bàn làm việc gỗ",
        "OFF-CAB-I": "Tủ hồ sơ",
        "OFF-PROJ-P": "Máy chiếu",
        "OFF-SCR-P": "Màn chiếu",
        "OFF-SHRED": "Máy hủy tài liệu",
        "OFF-BIND": "Máy đóng sách",
        "INK-CAN-12A": "Mực Canon 12A",
        "INK-CAN-35A": "Mực Canon 35A/85A",
        "INK-CAN-051": "Mực Canon 051/054",
        "INK-BRO-2385": "Mực Brother TN-2385",
        "INK-BRO-B022": "Mực Brother TN-B022",
        "INK-BRO-1010": "Mực Brother TN-1010",
        "INK-HP-17A": "Mực HP 17A/107A",
        "INK-EPS-673": "Mực Epson 673",
        "INK-EPS-003": "Mực Epson 003",
        "INK-CAN-71": "Mực Canon GI-71",
        "INK-RIB-80": "Ruy băng",
        "WST-DRUM": "Trống máy in (Drum)",
        "WST-ROLL": "Trục sấy / Rulo",
        "WST-CHIP": "Chip mực",
        "WST-PAPER-A4": "Giấy in A4",
        "WST-PAPER-BILL": "Giấy nhiệt BILL",
        "WST-INK-REFILL": "Mực nạp / Mực lẻ",
        "CPU-INT-I3": "CPU Intel Core i3",
        "CPU-INT-I5": "CPU Intel Core i5",
        "CPU-INT-I7": "CPU Intel Core i7",
        "CPU-AMD-RY": "CPU AMD Ryzen",
        "MB-H610": "Mainboard H610",
        "MB-B760": "Mainboard B760",
        "MB-Z790": "Mainboard Z790",
        "RAM-8G-D4": "RAM 8GB DDR4",
        "RAM-16G-D4": "RAM 16GB DDR4",
        "RAM-D5": "RAM DDR5",
        "VGA-GTX-16": "VGA GTX 16xx",
        "VGA-RTX-30": "VGA RTX 30xx",
        "VGA-RTX-40": "VGA RTX 40xx",
        "UPS-SAN-500": "UPS Santak",
        "UPS-APC-PRO": "UPS APC",
        "SSD-128-256": "SSD 128-256GB",
        "SSD-480-512": "SSD 480-512GB",
        "SSD-1T-2T": "SSD 1TB-2TB",
        "HDD-1TB": "HDD 1TB",
        "HDD-2TB-4TB": "HDD 2-4TB",
        "HDD-EXT": "HDD Di động",
        "USB-32G": "USB 32GB",
        "USB-64-128": "USB 64-128GB",
        "SD-CARD": "Thẻ nhớ",
        "CASE-OFF": "Vỏ máy văn phòng",
        "CASE-GAM": "Vỏ máy Gaming",
        "PSU-OFF": "Nguồn văn phòng",
        "PSU-GAM": "Nguồn thực 500W+",
        "COOL-FAN": "Quạt tản nhiệt",
        "COOL-AIO": "Tản nhiệt nước",
        "MS-LOGI": "Chuột Logitech",
        "MS-GAM": "Chuột Gaming",
        "MS-RAPO": "Chuột Rapoo",
        "KB-OFF": "Bàn phím văn phòng",
        "KB-MECH": "Bàn phím cơ",
        "HDSET-OFF": "Tai nghe",
        "SPK-2.0": "Loa máy tính",
        "HUB-USB": "Hub USB/Type-C",
        "CAM-WC": "Webcam",
        "NET-WF-HOME": "Router WiFi gia đình",
        "NET-WF-DUAL": "Router WiFi Dual Band",
        "NET-WF-MESH": "Hệ thống WiFi Mesh",
        "NET-WF-PRO": "WiFi Chuyên dụng",
        "NET-SW-OFF": "Switch văn phòng",
        "NET-SW-RACK": "Switch Rackmount",
        "NET-SW-POE": "Switch PoE",
        "NET-CARD": "Card mạng / USB WiFi",
        "NET-4G-W": "Bộ phát WiFi 4G",
        "NET-CAB-NET": "Tủ mạng / Rack",
        "CAM-WIFI-2M": "Camera WiFi 2MP",
        "CAM-WIFI-4M": "Camera WiFi 4MP",
        "CAM-IP-DOME": "Camera IP Dome",
        "CAM-IP-BUL": "Camera IP Bullet",
        "CAM-DVR-4C": "Đầu ghi 4 kênh",
        "CAM-DVR-8C": "Đầu ghi 8 kênh",
        "CAM-DVR-16": "Đầu ghi 16 kênh",
        "CAM-ACC-BS": "Chân đế Camera",
        "CAB-LAN-C5": "Cáp mạng Cat5e",
        "CAB-LAN-C6": "Cáp mạng Cat6",
        "CAB-HDMI-3": "Cáp HDMI ngắn",
        "CAB-HDMI-15": "Cáp HDMI dài",
        "CAB-VGA-DP": "Cáp VGA/DP/DVI",
        "VT-RJ45": "Đầu bấm mạng",
        "VT-RJ11": "Hạt điện thoại",
        "RADIO-W": "Bộ đàm",
        "TEL-PHONE": "Điện thoại bàn",
        "VT-TOOLS": "Công cụ thi công",
        "CAB-LAN-LOW": "Dây mạng lẻ",
        "WST-BATTERY": "Pin thiết bị",
        "SW-WIN-PRO": "Bản quyền Windows",
        "SW-OFF-365": "Bản quyền Office",
        "SW-AV-KAS": "Dịệt Virus",
        "SRV-INSTALL": "Phí lắp đặt",
        "SRV-MAINT": "Phí bảo trì",
        "OTH-GEN-LOW": "Hàng lẻ <50k",
        "OTH-GEN-MID": "Hàng hóa 50k-500k",
        "OTH-GEN-HIGH": "Hàng hóa 500k-5M",
        "OTH-PREMIUM": "Hàng giá trị cao >5M",
        "OTH-GEN": "Hàng hóa khác",
    }

def map_item(ten_hang):
    if not ten_hang: return "OTH-GEN"
    h = str(ten_hang).lower()
    if re.search(r'phí|lãi|vay|huy động|bảo hiểm|cước|quảng cáo|tiền điện|tiền nước|vận chuyển|thuê|sửa chữa|mặt bằng|triển khai', h):
        return "SERVICE-OUT"
    if re.search(r'latitude', h): return "PC-DELL-LAT"
    if re.search(r'vostro', h): return "PC-DELL-VOS"
    if re.search(r'inspiron', h): return "PC-DELL-INS"
    if re.search(r'xps', h): return "PC-DELL-XPS"
    if re.search(r'vivobook', h): return "PC-ASU-VIVO"
    if re.search(r'zenbook', h): return "PC-ASU-ZEN"
    if re.search(r'expertbook', h): return "PC-ASU-EXP"
    if re.search(r'rog|strix|zephyrus|tuf', h): return "PC-ASU-ROG"
    if re.search(r'thinkpad|t14|x1\s*carbon', h): return "PC-LEN-TP"
    if re.search(r'ideapad|slim\s*3|slim\s*5', h): return "PC-LEN-IP"
    if re.search(r'msi.*modern', h): return "PC-MSI-MOD"
    if re.search(r'katana|bravo', h): return "PC-MSI-GF"
    if re.search(r'pavilion', h): return "PC-HP-PAV"
    if re.search(r'probook', h): return "PC-HP-PRO"
    if re.search(r'elitebook', h): return "PC-HP-EL"
    if re.search(r'aspire', h): return "PC-ACER-ASP"
    if re.search(r'nitro', h): return "PC-ACER-NIT"
    if re.search(r'macbook.*air', h): return "PC-MAC-AIR"
    if re.search(r'macbook.*pro', h): return "PC-MAC-PRO"
    if re.search(r'ipad', h): return "PC-TAB-IPAD"
    if re.search(r'galaxy.*tab', h): return "PC-TAB-SAM"
    if re.search(r'pc.*i3|máy\s*tính.*i3|bộ\s*máy.*i3', h): return "PC-SYS-I3"
    if re.search(r'pc.*i5|máy\s*tính.*i5|bộ\s*máy.*i5', h): return "PC-SYS-I5"
    if re.search(r'pc.*i7|máy\s*tính.*i7|bộ\s*máy.*i7', h): return "PC-SYS-I7"
    if re.search(r'optiplex|proone|aio', h): return "PC-AIO-DELL"
    if re.search(r'server|xeon|workstation', h): return "PC-WS-XEON"
    if re.search(r'nuc|mini\s*pc', h): return "PC-MINI"
    if re.search(r'samsung.*19|samsung.*20', h): return "LCD-SAM-19"
    if re.search(r'samsung.*24|samsung.*27', h): return "LCD-SAM-24"
    if re.search(r'dell.*22|e22|p22', h): return "LCD-DELL-22"
    if re.search(r'dell.*24|u24|p24|e24', h): return "LCD-DELL-24"
    if re.search(r'dell.*27|u27|p27', h): return "LCD-DELL-27"
    if re.search(r'va24|vz24|vg24', h): return "LCD-ASU-24"
    if re.search(r'lg\s*24|lg\s*27', h): return "LCD-LG-24"
    if re.search(r'viewsonic', h): return "LCD-VIEW-24"
    if re.search(r'hp.*laserjet|m15|m102|m404', h): return "PRN-HP-LJ"
    if re.search(r'brother.*hl|2321|2361', h): return "PRN-BRO-HL"
    if re.search(r'epson.*l3110|l3210|l805|l1800', h): return "PRN-EPS-L"
    if re.search(r'ssd.*120|ssd.*128|ssd.*240|ssd.*256', h): return "SSD-128-256"
    if re.search(r'ssd.*480|ssd.*500|ssd.*512', h): return "SSD-480-512"
    if re.search(r'ssd.*1t|ssd.*2t', h): return "SSD-1T-2T"
    if re.search(r'hdd.*1tb', h): return "HDD-1TB"
    if re.search(r'hdd.*2tb|hdd.*4tb|skyhawk|purple', h): return "HDD-2TB-4TB"
    if re.search(r'hdd.*di.*động|ssd.*di.*động', h): return "HDD-EXT"
    if re.search(r'ram.*8g|8gb.*ram|ddr4.*8g', h): return "RAM-8G-D4"
    if re.search(r'ram.*16g|16gb.*ram|ddr4.*16g', h): return "RAM-16G-D4"
    if re.search(r'ram.*ddr5', h): return "RAM-D5"
    if re.search(r'i3-10100|i3-12100|i3-13100', h): return "CPU-INT-I3"
    if re.search(r'i5-10400|i5-11400|i5-12400|i5-13400', h): return "CPU-INT-I5"
    if re.search(r'i7-10700|i7-12700|i7-13700', h): return "CPU-INT-I7"
    if re.search(r'ryzen|r3-|r5-|r7-|r9-', h): return "CPU-AMD-RY"
    if re.search(r'h610', h): return "MB-H610"
    if re.search(r'b760|b660', h): return "MB-B760"
    if re.search(r'z790|z690', h): return "MB-Z790"
    if re.search(r'gtx.*1650|1660', h): return "VGA-GTX-16"
    if re.search(r'rtx.*3050|3060', h): return "VGA-RTX-30"
    if re.search(r'rtx.*4060|4070|4080', h): return "VGA-RTX-40"
    if re.search(r'mực|cartridge|toner', h):
        if re.search(r'12a|303', h): return "INK-CAN-12A"
        if re.search(r'35a|85a|78a|325', h): return "INK-CAN-35A"
        if re.search(r'051|054|057', h): return "INK-CAN-051"
        if re.search(r'2385', h): return "INK-BRO-2385"
        if re.search(r'b022|b027', h): return "INK-BRO-B022"
        if re.search(r'1010', h): return "INK-BRO-1010"
        if re.search(r'17a|107a', h): return "INK-HP-17A"
        if re.search(r'673', h): return "INK-EPS-673"
        if re.search(r'003', h): return "INK-EPS-003"
        if re.search(r'71', h): return "INK-CAN-71"
    if re.search(r'canon.*lbp|lbp|2900|6030|6230|223dw', h): return "PRN-CAN-LBP"
    if re.search(r'canon.*mf|d1650|mf241|mf269', h): return "PRN-CAN-MF"
    if re.search(r'brother.*hl|2321|2351|2361|2366', h): return "PRN-BRO-HL"
    if re.search(r'brother.*dcp|b7535|t420|t720|t820', h): return "PRN-BRO-DCP"
    if re.search(r'laserjet|m102|m404|4003', h) or (re.search(r'hp', h) and re.search(r'in', h)): return "PRN-HP-LJ"
    if re.search(r'epson\s*l|l3110|l3210|l805|l1800', h): return "PRN-EPS-L"
    if re.search(r'quét|scan.*canon', h): return "SCN-CAN-CANO"
    if re.search(r'xprinter|k80|in\s*hóa\s*đơn|xp-', h): return "PRN-POS-80"
    if re.search(r'mã\s*vạch|barcode|honeywell|zebra', h): return "SCN-BAR"
    if re.search(r'bảo.*trì|sửa.*chữa', h): return "SRV-MAINT"
    return "OTH-GEN"

# ============================================================
# DATA RETRIEVAL
# ============================================================
def fetch_raw_data(engine, year):
    print(f"--- FETCHING DATA FOR {year} ---", flush=True)
    query_list = text("""
        SELECT "idServer", shdon, loaihd, tthai, tgtcthue, nbten, nmten,
               tdlap AT TIME ZONE 'UTC' AT TIME ZONE 'Asia/Ho_Chi_Minh' as tdlap_ict
        FROM ext_listhoadon
        WHERE "congtyId" = :cid
          AND tthai IN ('1','2','4','5')
          AND (tdlap AT TIME ZONE 'UTC' AT TIME ZONE 'Asia/Ho_Chi_Minh') >= :start
          AND (tdlap AT TIME ZONE 'UTC' AT TIME ZONE 'Asia/Ho_Chi_Minh') < :end
    """)
    with engine.connect() as conn:
        df_list = pd.read_sql(query_list, conn, params={
            'cid': COMPANY_ID,
            'start': f'{year}-01-01',
            'end': f'{year+1}-01-01'
        })
    
    if df_list.empty: return pd.DataFrame(), pd.DataFrame()

    ids = df_list['idServer'].tolist()
    all_details = []
    batch_size = 200
    for i in range(0, len(ids), batch_size):
        batch = ids[i:i+batch_size]
        query_detail = text(f"""
            SELECT "idhdonServer", ten, sluong, dgia, thtien
            FROM ext_detailhoadon
            WHERE "idhdonServer" IN :batch
        """)
        with engine.connect() as conn:
            df_batch = pd.read_sql(query_detail, conn, params={'batch': tuple(batch)})
        all_details.append(df_batch)
    
    df_detail = pd.concat(all_details, ignore_index=True)
    return df_list, df_detail

def process_year(year, engine, prev_opening_balance=None, skip_list=None):
    df_list, df_detail = fetch_raw_data(engine, year)
    if df_list.empty: return None

    # 1. Calculate Reported Accounting Totals (Header Level)
    res_data = {
        'year': year,
        'reported_sums': {m: {'banra': 0, 'muavao': 0} for m in range(1, 13)},
        'hoadon_records': []
    }
    df_list['ict'] = pd.to_datetime(df_list['tdlap_ict'])
    
    # Pre-calculated sums using mask (More reliable than manual loop with NaN)
    for ltype in ['banra', 'muavao']:
        mask = (df_list['loaihd'] == ltype)
        if year == 2023:
            skip_vals = SALES_SKIP_2023 if ltype == 'banra' else PURCH_SKIP_2023
            mask &= ~df_list['shdon'].isin(skip_vals)
        elif skip_list:
            mask &= ~df_list['shdon'].isin(skip_list)
            
        group_m = df_list[mask].groupby(df_list['ict'].dt.month)['tgtcthue'].sum().to_dict()
        for m, val in group_m.items():
            res_data['reported_sums'][m][ltype] = val

    print(f"[{year}] Final Reported Sales (Header-sum): {sum(v['banra'] for v in res_data['reported_sums'].values()):,.0f}")
    print(f"[{year}] Final Reported Purch (Header-sum): {sum(v['muavao'] for v in res_data['reported_sums'].values()):,.0f}")

    # 2. Join with Details for Internal Inventory Logic
    df = df_detail.merge(df_list, left_on='idhdonServer', right_on='idServer')
    df['ict'] = pd.to_datetime(df['tdlap_ict'])
    df['group'] = df['ten'].apply(map_item)
    category_names = get_category_mapping()
    
    all_groups = sorted(set(list(df['group'].unique()) + (list(prev_opening_balance.keys()) if prev_opening_balance else [])))
    if "SERVICE-OUT" in all_groups: all_groups.remove("SERVICE-OUT")
    res_data['all_groups'] = all_groups
    res_data['monthly_data'] = {m: {g: {
            'ton_dau_sl': 0, 'ton_dau_val': 0,
            'nhap_sl': 0, 'nhap_val': 0,
            'xuat_sl': 0, 'xuat_val_ban': 0,
            'wap': 0, 'xuat_val_cogs': 0,
            'ton_cuoi_sl': 0, 'ton_cuoi_val': 0
        } for g in all_groups} for m in range(1, 13)}

    all_items = df.to_dict('records')
    # Opening carry forward
    for g in all_groups:
        if prev_opening_balance and g in prev_opening_balance:
            res_data['monthly_data'][1][g]['ton_dau_sl'] = prev_opening_balance[g]['qty']
            res_data['monthly_data'][1][g]['ton_dau_val'] = prev_opening_balance[g]['val']

    all_items.sort(key=lambda x: (x['ict'], x['loaihd'] == 'muavao'))
    current_qty = {g: (prev_opening_balance[g]['qty'] if prev_opening_balance and g in prev_opening_balance else 0) for g in all_groups}
    current_val = {g: (prev_opening_balance[g]['val'] if prev_opening_balance and g in prev_opening_balance else 0) for g in all_groups}
    
    bank_keywords = ['ngân hàng', 'bank', 'bidv', 'vietin', 'agri', 'sacom', 'vcb', 'acb']

    for item in all_items:
        m = item['ict'].month
        sh = str(item['shdon'])
        cat = item['group']
        partner = str(item['nbten'] if item['loaihd'] == 'muavao' else item['nmten'])
        is_bank = any(k in partner.lower() for k in bank_keywords)
        
        should_skip_inv = (skip_list and sh in skip_list)
        if year == 2023:
            if item['loaihd'] == 'banra' and sh in SALES_SKIP_2023: should_skip_inv = True
            if item['loaihd'] == 'muavao' and sh in PURCH_SKIP_2023: should_skip_inv = True
            
        res_data['hoadon_records'].append({
            'Tháng': m, 'Ngày': item['ict'].strftime("%Y-%m-%d"), 'Số HĐ': sh,
            'Loại': item['loaihd'], 'Đối tác': partner, 'Nội dung': item['ten'],
            'Nhóm': cat, 'SL': item['sluong'], 'Đơn giá': item['dgia'], 'Thành tiền': item['thtien']
        })

        if cat == "SERVICE-OUT" or should_skip_inv or is_bank: continue

        m_data = res_data['monthly_data'][m][cat]
        if item['loaihd'] == 'muavao':
            m_data['nhap_sl'] += item['sluong']
            m_data['nhap_val'] += item['thtien']
            current_qty[cat] += item['sluong']
            current_val[cat] += item['thtien']
        else:
            m_data['xuat_sl'] += item['sluong']
            m_data['xuat_val_ban'] += item['thtien']
            wap = current_val[cat] / current_qty[cat] if current_qty[cat] > 0 else 0
            cogs = item['sluong'] * wap
            m_data['xuat_val_cogs'] += cogs
            m_data['wap'] = wap
            current_qty[cat] -= item['sluong']
            current_val[cat] -= cogs

    for m in range(1, 13):
        for g in all_groups:
            d = res_data['monthly_data'][m][g]
            d['ton_cuoi_sl'] = d['ton_dau_sl'] + d['nhap_sl'] - d['xuat_sl']
            d['ton_cuoi_val'] = d['ton_dau_val'] + d['nhap_val'] - d['xuat_val_cogs']
            if m < 12:
                next_d = res_data['monthly_data'][m+1][g]
                next_d['ton_dau_sl'] = d['ton_cuoi_sl']
                next_d['ton_dau_val'] = d['ton_cuoi_val']
                
    res_data['closing_balance'] = {g: {'qty': current_qty[g], 'val': current_val[g]} for g in all_groups}
    return res_data

def save_excel(year_results, output_path):
    wb = Workbook()
    wb.remove(wb.active)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2F5496")
    total_font = Font(bold=True)
    title_font = Font(bold=True, size=14)
    border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    category_names = get_category_mapping()
    year = year_results['year']

    ws_rep = wb.create_sheet("Đối Chiếu Mục Tiêu")
    ws_rep.cell(1, 1, f"ĐỐI CHIẾU DOANH THU & CHI PHÍ NĂM {year} (HÓA ĐƠN)").font = title_font
    rep_cols = ['Tháng', 'Tổng Bán Ra (Hóa Đơn)', 'Tổng Mua Vào (Hóa Đơn)']
    for c, h in enumerate(rep_cols, 1):
        cell = ws_rep.cell(3, c, h); cell.font = header_font; cell.fill = header_fill; cell.border = border
    for m in range(1, 13):
        r = m + 3
        ws_rep.cell(r, 1, f"Tháng {m}").border = border
        ws_rep.cell(r, 2, year_results['reported_sums'][m]['banra']).number_format = '#,##0'; ws_rep.cell(r, 2).border = border
        ws_rep.cell(r, 3, year_results['reported_sums'][m]['muavao']).number_format = '#,##0'; ws_rep.cell(r, 3).border = border
    ws_rep.cell(16, 1, "TỔNG CỘNG").font = total_font
    ws_rep.cell(16, 2, f"=SUM(B4:B15)").number_format = '#,##0'; ws_rep.cell(16, 2).font = total_font
    ws_rep.cell(16, 3, f"=SUM(C4:C15)").number_format = '#,##0'; ws_rep.cell(16, 3).font = total_font

    ws_all = wb.create_sheet("xnt12thang")
    sum_headers = ['STT', 'Mã Nhóm', 'Tên Nhóm', 'Tồn Đầu Năm', 'Tổng Nhập', 'Tổng Xuất (Bán)', 'Giá Vốn', 'Tồn Cuối Năm']
    for m in range(1, 13): sum_headers.extend([f"Nhập T{m}", f"Xuất T{m}"])
    for c, h in enumerate(sum_headers, 1): cell = ws_all.cell(1, c, h); cell.font = header_font; cell.fill = header_fill; cell.border = border
    row_idx = 2
    for idx, g in enumerate(year_results['all_groups'], 1):
        d12 = year_results['monthly_data'][12][g]
        ton_dau = year_results['monthly_data'][1][g]['ton_dau_val']
        t_nhap = sum(year_results['monthly_data'][m][g]['nhap_val'] for m in range(1, 13))
        t_xuat = sum(year_results['monthly_data'][m][g]['xuat_val_ban'] for m in range(1, 13))
        t_cogs = sum(year_results['monthly_data'][m][g]['xuat_val_cogs'] for m in range(1, 13))
        row = [idx, g, category_names.get(g, g), ton_dau, t_nhap, t_xuat, t_cogs, d12['ton_cuoi_val']]
        for m in range(1, 13): row.extend([year_results['monthly_data'][m][g]['nhap_val'], year_results['monthly_data'][m][g]['xuat_val_ban']])
        for c, v in enumerate(row, 1):
            cell = ws_all.cell(row_idx, c, v); cell.border = border
            cell.number_format = '#,##0' if c >= 4 else 'General'
        row_idx += 1

    for m in range(1, 13):
        ws = wb.create_sheet(f"Tháng {m}")
        headers = ['STT', 'Mã Nhóm', 'Tên Nhóm', 'Tồn Đầu (SL)', 'Tồn Đầu (VNĐ)', 'Nhập (SL)', 'Nhập (VNĐ)', 'Xuất (SL)', 'Xuất (VNĐ)', 'Giá Vốn (WAP)', 'Giá Vốn (VNĐ)', 'Tồn Cuối (SL)', 'Tồn Cuối (VNĐ)']
        for c, h in enumerate(headers, 1): cell = ws.cell(1, c, h); cell.font = header_font; cell.fill = header_fill; cell.border = border
        r_idx = 2
        for idx, g in enumerate(year_results['all_groups'], 1):
            d = year_results['monthly_data'][m][g]
            row = [idx, g, category_names.get(g, g), d['ton_dau_sl'], d['ton_dau_val'], d['nhap_sl'], d['nhap_val'], d['xuat_sl'], d['xuat_val_ban'], d['wap'], d['xuat_val_cogs'], d['ton_cuoi_sl'], d['ton_cuoi_val']]
            for c, v in enumerate(row, 1):
                cell = ws.cell(r_idx, c, v); cell.border = border
                cell.number_format = '#,##0' if c >= 4 else 'General'
            r_idx += 1

    ws_hd = wb.create_sheet("Hoadon")
    hd_headers = ['Tháng', 'Ngày', 'Số HĐ', 'Loại', 'Đối tác', 'Nội dung', 'Nhóm', 'SL', 'Đơn giá', 'Thành tiền']
    for c, h in enumerate(hd_headers, 1): cell = ws_hd.cell(1, c, h); cell.font = header_font; cell.fill = header_fill; cell.border = border
    for r, rec in enumerate(year_results['hoadon_records'], 2):
        for c, h in enumerate(hd_headers, 1):
            cell = ws_hd.cell(r, c, rec.get(h, "")); cell.border = border
            cell.number_format = '#,##0' if c >= 9 else 'General'

    wb.save(output_path)
    print(f"Excel saved: {output_path}", flush=True)

def main():
    engine = create_engine(DB_URI)
    years = [2023, 2024, 2025, 2026]
    prev_closing = {'OTH-GEN': {'qty': 20000, 'val': 20528682383.0}}
    for year in years:
        skip_path = f"/chikiet/kata2025/ragketoan/python/skip_lists/skip_list_{year}.json"
        skip_shdons = set()
        if os.path.exists(skip_path):
            with open(skip_path, 'r', encoding='utf-8') as f:
                data = json.load(f); entries = data.get('skip_entries', []) if isinstance(data, dict) else data
                for s in entries: skip_shdons.add(str(s.get('shdon', s) if isinstance(s, dict) else s))
        print(f"Processing year {year}...", flush=True)
        results = process_year(year, engine, prev_opening_balance=prev_closing, skip_list=skip_shdons)
        if results:
            save_excel(results, os.path.join(OUTPUT_DIR, f"XNT_HuyVu_{year}.xlsx"))
            prev_closing = results['closing_balance']
            print(f"Done {year}. Next Opening: {sum(d['val'] for d in prev_closing.values()):,.0f}")

if __name__ == "__main__":
    main()
